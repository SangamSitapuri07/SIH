#!/usr/bin/env python3
"""Native-Object scenes for the Slave Dynasty deck.

Every illustration is built from REAL PowerPoint objects — AutoShapes,
connectors and freeform silhouettes — and animated with PowerPoint's own
animation engine (entrances, motion paths, spins, pulses) via the PML
Animator in make_ppt.py.  No rasterised GIF frames: everything on the slide
stays selectable, recolourable and editable in PowerPoint.

Scene builders signature:  scene_xxx(slide, anim, L, T, W, H)
with (L, T) = top-left of the illustration box and (W, H) its size (EMU).
Inside, scenes are laid out on a normalised 3.5 x 4.4 inch canvas.
"""

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN

EMU = 914400
SW_IN, SH_IN = 13.333, 7.5          # slide size in inches (for motion paths)

CREAM2 = RGBColor(0xFF, 0xE9, 0xC9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BROWN = RGBColor(0x80, 0x34, 0x0D)
BROWN2 = RGBColor(0xA8, 0x55, 0x1E)
DARK = RGBColor(0x26, 0x1A, 0x10)
NIGHT = RGBColor(0x20, 0x28, 0x44)
GREEN2 = RGBColor(0x2E, 0x6B, 0x33)
GOLD = RGBColor(0xC8, 0x8A, 0x3C)
GRAY = RGBColor(0x7F, 0x7F, 0x7F)
BLUE = RGBColor(0x3B, 0x6E, 0xA5)
SILVER = RGBColor(0xCD, 0xC9, 0xBE)
TAN = RGBColor(0xC8, 0x9B, 0x6A)
TEDDY = RGBColor(0x9C, 0x6B, 0x3F)
TEDDY2 = RGBColor(0xD9, 0xB4, 0x8C)

HEAD_FONT = "Bernard MT Condensed"
BODY_FONT = "Times New Roman"


# ---------------------------------------------------------------- helpers
def _shp(s, kind, l, t, w, h, fill=None, line=None, lw=1.0, rot=0):
    sh = s.shapes.add_shape(kind, int(l), int(t), int(w), int(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    if rot:
        sh.rotation = rot
    return sh


def _txt(sh, lines):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Pt(2)
    tf.margin_top = tf.margin_bottom = Pt(1)
    for i, (t, sz, c, f, b) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = t
        r.font.size = Pt(sz); r.font.color.rgb = c
        r.font.name = f; r.font.bold = b
    return sh


def _free(s, verts, l, t, fill, line=None, lw=1.0):
    """Closed freeform silhouette; verts in local EMU offsets from (l, t)."""
    b = s.shapes.build_freeform(verts[0][0], verts[0][1], scale=1.0)
    b.add_line_segments(verts[1:], close=True)
    sh = b.convert_to_shape(origin_x=int(l), origin_y=int(t))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


def _scale_verts(verts, sx, sy, mirror=False, width=0.0):
    out = []
    for x, y in verts:
        xx = x * sx
        if mirror:
            xx = width - xx
        out.append((int(xx), int(y * sy)))
    return out


# galloping horse silhouette, drawn in a 1.2 x 0.93 inch box
HORSE_V = [(1.05, 0.18), (1.14, 0.30), (1.02, 0.36), (0.92, 0.44),
           (0.97, 0.62), (1.00, 0.90), (0.93, 0.90), (0.88, 0.64),
           (0.62, 0.68), (0.57, 0.92), (0.50, 0.92), (0.52, 0.64),
           (0.32, 0.60), (0.30, 0.88), (0.23, 0.88), (0.25, 0.56),
           (0.06, 0.52), (0.00, 0.34), (0.16, 0.42), (0.36, 0.38),
           (0.62, 0.36), (0.84, 0.20), (0.90, 0.06), (0.96, 0.16)]

# talwar silhouette, drawn in a 0.30 x 1.92 inch box (tip up)
SWORD_V = [(0.15, 0.00), (0.19, 0.25), (0.19, 1.35), (0.30, 1.38),
           (0.30, 1.46), (0.20, 1.46), (0.20, 1.76), (0.24, 1.84),
           (0.15, 1.92), (0.06, 1.84), (0.10, 1.76), (0.10, 1.46),
           (0.00, 1.46), (0.00, 1.38), (0.11, 1.35), (0.11, 0.25)]


def horse_shape(s, l, t, w, h, fill=DARK):
    v = _scale_verts(HORSE_V, w / 1.2, h / 0.93)
    return _free(s, v, l, t, fill)


def sword_shape(s, l, t, w, h, mirror=False, rot=0, fill=SILVER, line=BROWN):
    v = _scale_verts(SWORD_V, w / 0.30, h / 1.92, mirror=mirror, width=w)
    sh = _free(s, v, l, t, fill, line=line, lw=1.0)
    if rot:
        sh.rotation = rot
    return sh


def arrow_shape(s, l, t, w, h, fill=BROWN):
    ar = _shp(s, MSO_SHAPE.RIGHT_ARROW, l, t, w, h, fill=fill)
    ar.adjustments[0] = 0.55
    ar.adjustments[1] = 0.55
    return ar


class Canvas:
    """Normalised 3.5 x 4.4 inch drawing canvas inside the illustration box."""

    def __init__(self, s, L, T, W, H, cw=3.5, ch=4.4):
        self.s = s
        self.sx = W / cw            # EMU per canvas inch, horizontal
        self.sy = H / ch
        self.L, self.T = L, T

    def X(self, x): return self.L + x * self.sx
    def Y(self, y): return self.T + y * self.sy
    def W(self, w): return w * self.sx
    def H(self, h): return h * self.sy

    def shp(self, kind, x, y, w, h, **kw):
        return _shp(self.s, kind, self.X(x), self.Y(y), self.W(w), self.H(h), **kw)

    def free(self, verts, x, y, sx, sy, **kw):
        v = _scale_verts(verts, sx, sy)
        return _free(self.s, v, self.X(x), self.Y(y), **kw)

    def fx(self, inches): return inches / SW_IN      # motion-path fractions
    def fy(self, inches): return inches / SH_IN


def _cap(s, anim, L, T, W, text):
    from pptx.util import Inches as _I
    tb = s.shapes.add_textbox(int(L - _I(0.3)), int(T + W * 0 + 0), 1, 1)  # placeholder
    return tb


# ================================================================= scenes
def scene_intro(s, anim, L, T, W, H):
    """From chains to crowns: a broken chain, an arrow, a crown."""
    c = Canvas(s, L, T, W, H)
    links = []
    for i in range(3):
        lk = c.shp(MSO_SHAPE.OVAL, 0.55 + i * 0.42, 3.15 - i * 0.28, 0.55, 0.34,
                   fill=None, line=DARK, lw=3.0, rot=(-24 if i % 2 == 0 else 24))
        links.append(lk)
    broke = c.shp(MSO_SHAPE.OVAL, 1.85, 2.28, 0.55, 0.34, fill=None, line=DARK,
                  lw=3.0, rot=64)
    spark = c.shp(MSO_SHAPE.STAR_8_POINT, 2.28, 2.02, 0.6, 0.6, fill=GOLD, line=BROWN, lw=1.25)
    arrow = c.shp(MSO_SHAPE.UP_ARROW, 1.42, 1.15, 0.62, 0.95, fill=GREEN2)
    band = c.shp(MSO_SHAPE.RECTANGLE, 0.95, 0.72, 1.6, 0.34, fill=GOLD, line=BROWN, lw=1.25)
    pts = [c.shp(MSO_SHAPE.ISOSCELES_TRIANGLE, 0.95 + i * 0.53, 0.30, 0.5, 0.45,
                 fill=GOLD, line=BROWN, lw=1.25) for i in range(3)]
    gem = c.shp(MSO_SHAPE.OVAL, 1.62, 0.79, 0.2, 0.2, fill=BROWN)
    lab = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 0.55, 3.72, 2.4, 0.5, fill=BROWN)
    _txt(lab, [("slave \u2192 sultan", 13, CREAM2, HEAD_FONT, False)])
    effs, d = [], 0
    for sp in links + [broke]:
        effs.append({"spid": sp.shape_id, "kind": "fade", "dur": 300, "delay": d}); d += 260
    effs.append({"spid": spark.shape_id, "kind": "zoom", "dur": 300, "delay": d})
    effs.append({"spid": arrow.shape_id, "kind": "wipeup", "dur": 450, "delay": d + 250})
    for sp in [band] + pts + [gem]:
        effs.append({"spid": sp.shape_id, "kind": "zoom", "dur": 350, "delay": d + 650})
    effs.append({"spid": lab.shape_id, "kind": "fade", "dur": 400, "delay": d + 1000})
    anim.add_group(effs, click=False)
    anim.add_group([{"spid": gem.shape_id, "kind": "pulse", "dur": 800,
                     "repeat": "indefinite"}], click=False)
    for sp in links + [broke, spark, arrow, band, gem, lab] + pts:
        anim.build_shape(sp.shape_id)


def scene_market(s, anim, L, T, W, H):
    """Slave market at Bukhara: stall, boy, price tag, raining coins."""
    c = Canvas(s, L, T, W, H)
    awn = c.shp(MSO_SHAPE.RECTANGLE, 0.45, 0.55, 2.6, 0.42, fill=BROWN2, line=GOLD, lw=1.25)
    leg1 = c.shp(MSO_SHAPE.RECTANGLE, 0.5, 0.97, 0.12, 2.1, fill=BROWN)
    leg2 = c.shp(MSO_SHAPE.RECTANGLE, 2.88, 0.97, 0.12, 2.1, fill=BROWN)
    plat = c.shp(MSO_SHAPE.RECTANGLE, 0.45, 3.05, 2.6, 0.28, fill=BROWN, line=GOLD, lw=1.0)
    head = c.shp(MSO_SHAPE.OVAL, 1.52, 1.62, 0.46, 0.46, fill=TAN)
    body = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 1.42, 2.10, 0.66, 0.95, fill=BROWN2)
    tag = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 0.62, 1.35, 1.05, 0.72, fill=CREAM2, line=BROWN, lw=1.25)
    _txt(tag, [("SOLD", 12, BROWN, HEAD_FONT, True),
               ("100 dirhams", 9.5, DARK, BODY_FONT, False)])
    coins = [c.shp(MSO_SHAPE.OVAL, 0.75 + i * 0.55, 0.15, 0.3, 0.3, fill=GOLD, line=BROWN, lw=1.0)
             for i in range(4)]
    anim.add_group([{"spid": awn.shape_id, "kind": "wipe", "dur": 400},
                    {"spid": leg1.shape_id, "kind": "fade", "dur": 250, "delay": 300},
                    {"spid": leg2.shape_id, "kind": "fade", "dur": 250, "delay": 300},
                    {"spid": plat.shape_id, "kind": "wipe", "dur": 350, "delay": 450},
                    {"spid": body.shape_id, "kind": "zoom", "dur": 450, "delay": 700},
                    {"spid": head.shape_id, "kind": "zoom", "dur": 450, "delay": 700},
                    {"spid": tag.shape_id, "kind": "fade", "dur": 400, "delay": 1150}],
                   click=False)
    for i, cn in enumerate(coins):
        anim.add_group([{"spid": cn.shape_id, "kind": "motion",
                         "path": f"M 0 0 L 0 {c.fy(2.6):.3f} E",
                         "dur": 1500 + i * 200, "delay": i * 350,
                         "repeat": "indefinite"}], click=False)
    for sp in [awn, leg1, leg2, plat, head, body, tag] + coins:
        anim.build_shape(sp.shape_id)


def scene_ghori(s, anim, L, T, W, H):
    """Damyak, March 1206: night camp, three daggers fly in."""
    c = Canvas(s, L, T, W, H)
    night = c.shp(MSO_SHAPE.RECTANGLE, 0.15, 0.2, 3.2, 3.5, fill=NIGHT)
    moon = c.shp(MSO_SHAPE.OVAL, 2.55, 0.5, 0.55, 0.55, fill=CREAM2)
    tent = c.shp(MSO_SHAPE.ISOSCELES_TRIANGLE, 0.85, 2.15, 1.8, 1.35, fill=BROWN2, line=GOLD, lw=1.0)
    gate = c.shp(MSO_SHAPE.RECTANGLE, 1.55, 2.85, 0.4, 0.65, fill=DARK)
    daggers = [sword_shape(s, c.X(x), c.Y(y), c.W(0.16), c.H(1.0), mirror=m, rot=r,
                           fill=SILVER, line=GOLD)
               for (x, y, m, r) in ((0.35, 0.55, False, 35), (2.85, 0.75, True, -35),
                                    (1.6, 0.35, False, 0))]
    star = c.shp(MSO_SHAPE.STAR_8_POINT, 1.42, 1.55, 0.7, 0.7, fill=GOLD, line=BROWN, lw=1.25)
    lab = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 3.85, 2.1, 0.45, fill=BROWN)
    _txt(lab, [("no heir \u2014 only swords", 11.5, CREAM2, HEAD_FONT, False)])
    anim.add_group([{"spid": night.shape_id, "kind": "fade", "dur": 400},
                    {"spid": moon.shape_id, "kind": "fade", "dur": 400, "delay": 250},
                    {"spid": tent.shape_id, "kind": "wipeup", "dur": 500, "delay": 450},
                    {"spid": gate.shape_id, "kind": "fade", "dur": 300, "delay": 800}],
                   click=False)
    anim.add_group([
        {"spid": daggers[0].shape_id, "kind": "fly", "path": "M -0.12 -0.1 L 0 0 E",
         "dur": 500, "swing": 30},
        {"spid": daggers[1].shape_id, "kind": "fly", "path": "M 0.12 -0.1 L 0 0 E",
         "dur": 500, "delay": 250, "swing": -30},
        {"spid": daggers[2].shape_id, "kind": "fly", "path": "M 0 -0.14 L 0 0 E",
         "dur": 500, "delay": 500},
        {"spid": star.shape_id, "kind": "zoom", "dur": 300, "delay": 950},
        {"spid": lab.shape_id, "kind": "fade", "dur": 400, "delay": 1250}], click=False)
    anim.add_group([{"spid": star.shape_id, "kind": "pulse", "dur": 700,
                     "repeat": "indefinite"}], click=False)
    for sp in [night, moon, tent, gate, star, lab] + daggers:
        anim.build_shape(sp.shape_id)


def scene_aibak(s, anim, L, T, W, H):
    """Lakh Baksh: gold rains from the throne."""
    c = Canvas(s, L, T, W, H)
    back = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 1.05, 1.15, 1.4, 1.7, fill=BROWN, line=GOLD, lw=1.5)
    seat = c.shp(MSO_SHAPE.RECTANGLE, 0.85, 2.7, 1.8, 0.45, fill=BROWN2, line=GOLD, lw=1.0)
    base = c.shp(MSO_SHAPE.RECTANGLE, 0.65, 3.15, 2.2, 0.3, fill=BROWN)
    crown = c.shp(MSO_SHAPE.STAR_5_POINT, 1.45, 0.55, 0.6, 0.6, fill=GOLD, line=BROWN, lw=1.25)
    coins = [c.shp(MSO_SHAPE.OVAL, 0.55 + i * 0.62, 0.25, 0.32, 0.32, fill=GOLD, line=BROWN, lw=1.0)
             for i in range(5)]
    lab = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 3.7, 2.3, 0.5, fill=BROWN)
    _txt(lab, [("LAKH BAKSH \u2014 giver of lakhs", 12, CREAM2, HEAD_FONT, False)])
    anim.add_group([{"spid": back.shape_id, "kind": "wipeup", "dur": 500},
                    {"spid": seat.shape_id, "kind": "wipe", "dur": 350, "delay": 400},
                    {"spid": base.shape_id, "kind": "wipe", "dur": 350, "delay": 550},
                    {"spid": crown.shape_id, "kind": "zoom", "dur": 400, "delay": 800},
                    {"spid": lab.shape_id, "kind": "fade", "dur": 400, "delay": 1100}],
                   click=False)
    for i, cn in enumerate(coins):
        anim.add_group([{"spid": cn.shape_id, "kind": "motion",
                         "path": f"M 0 0 L {c.fx(0.12):.3f} {c.fy(3.2):.3f} E",
                         "dur": 1700 + i * 180, "delay": i * 300,
                         "repeat": "indefinite"}], click=False)
    anim.add_group([{"spid": crown.shape_id, "kind": "pulse", "dur": 900,
                     "repeat": "indefinite"}], click=False)
    for sp in [back, seat, base, crown, lab] + coins:
        anim.build_shape(sp.shape_id)


def scene_iltutmish(s, anim, L, T, W, H):
    """1229: the Caliph's robe arrives from Baghdad; caravan loops."""
    c = Canvas(s, L, T, W, H)
    road = c.shp(MSO_SHAPE.RECTANGLE, 0.15, 3.35, 3.2, 0.1, fill=GOLD)
    scroll = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 1.0, 1.0, 1.6, 1.05, fill=CREAM2, line=BROWN, lw=1.5)
    _txt(scroll, [("robe of honour", 12, BROWN, HEAD_FONT, True),
                  ("from Baghdad", 10.5, DARK, BODY_FONT, False),
                  ("1229", 10.5, BROWN2, BODY_FONT, True)])
    seal = c.shp(MSO_SHAPE.OVAL, 2.35, 1.85, 0.5, 0.5, fill=DARK, line=GOLD, lw=1.5)
    star = c.shp(MSO_SHAPE.STAR_5_POINT, 2.47, 1.97, 0.26, 0.26, fill=GOLD)
    dots = [c.shp(MSO_SHAPE.OVAL, 0.2 + i * 0.5, 3.1, 0.26, 0.26, fill=BROWN)
            for i in range(3)]
    lab = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 0.5, 3.75, 2.5, 0.5, fill=BROWN)
    _txt(lab, [("a lawful sultanate at last", 12, CREAM2, HEAD_FONT, False)])
    anim.add_group([{"spid": road.shape_id, "kind": "wipe", "dur": 500},
                    {"spid": scroll.shape_id, "kind": "fly",
                     "path": f"M -{c.fx(2.2):.3f} -0.04 L 0 0 E", "dur": 1100, "delay": 300},
                    {"spid": seal.shape_id, "kind": "zoom", "dur": 350, "delay": 1400},
                    {"spid": star.shape_id, "kind": "zoom", "dur": 300, "delay": 1600},
                    {"spid": lab.shape_id, "kind": "fade", "dur": 400, "delay": 1850}],
                   click=False)
    for i, d in enumerate(dots):
        anim.add_group([{"spid": d.shape_id, "kind": "motion",
                         "path": f"M 0 0 L {c.fx(3.0):.3f} 0 E",
                         "dur": 4200, "delay": i * 700, "repeat": "indefinite"}],
                       click=False)
    for sp in [road, scroll, seal, star, lab] + dots:
        anim.build_shape(sp.shape_id)


def scene_mongol(s, anim, L, T, W, H):
    """1221: the horde sways at the Indus and turns away."""
    c = Canvas(s, L, T, W, H)
    river = c.shp(MSO_SHAPE.RECTANGLE, 1.62, 0.35, 0.24, 3.1, fill=BLUE)
    rlab = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 1.42, 3.5, 0.65, 0.36, fill=BLUE)
    _txt(rlab, [("INDUS", 9.5, WHITE, BODY_FONT, True)])
    horde = [c.shp(MSO_SHAPE.OVAL, 0.25 + i * 0.42, 1.3 + (0.35 if i % 2 else 0), 0.62, 0.5,
                   fill=DARK) for i in range(3)]
    khan = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 0.3, 0.55, 1.0, 0.55, fill=DARK, line=GOLD, lw=1.0)
    _txt(khan, [("CHINGIZ", 11, CREAM2, HEAD_FONT, False)])
    delhi = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 2.25, 1.5, 1.0, 0.75, fill=BROWN, line=GOLD, lw=1.25)
    _txt(delhi, [("DELHI", 12, CREAM2, HEAD_FONT, False),
                 ("refuses", 9.5, WHITE, BODY_FONT, True)])
    arrows = [arrow_shape(s, c.X(0.15), c.Y(2.6 + i * 0.3), c.W(0.7), c.H(0.16), fill=DARK)
              for i in range(3)]
    anim.add_group([{"spid": river.shape_id, "kind": "wipeup", "dur": 500},
                    {"spid": rlab.shape_id, "kind": "fade", "dur": 300, "delay": 400},
                    {"spid": khan.shape_id, "kind": "fade", "dur": 400, "delay": 550},
                    {"spid": horde[0].shape_id, "kind": "fade", "dur": 300, "delay": 750},
                    {"spid": horde[1].shape_id, "kind": "fade", "dur": 300, "delay": 900},
                    {"spid": horde[2].shape_id, "kind": "fade", "dur": 300, "delay": 1050},
                    {"spid": delhi.shape_id, "kind": "zoom", "dur": 450, "delay": 1250}],
                   click=False)
    anim.add_group([{"spid": h.shape_id, "kind": "motion",
                     "path": f"M 0 0 L {c.fx(0.55):.3f} 0 L 0 0 E",
                     "dur": 5200, "delay": i * 250, "repeat": "indefinite"}
                    for i, h in enumerate(horde)], click=False)
    for i, ar in enumerate(arrows):
        anim.add_group([{"spid": ar.shape_id, "kind": "motion",
                         "path": f"M 0 0 L {c.fx(1.15):.3f} 0 E", "dur": 1600 + i * 250,
                         "delay": i * 400, "repeat": "indefinite"}], click=False)
    anim.add_group([{"spid": delhi.shape_id, "kind": "pulse", "dur": 900,
                     "repeat": "indefinite"}], click=False)
    for sp in [river, rlab, khan, delhi] + horde + arrows:
        anim.build_shape(sp.shape_id)


def scene_razia(s, anim, L, T, W, H):
    """Razia rides at the head of her army — looping gallop."""
    c = Canvas(s, L, T, W, H)
    ground = c.shp(MSO_SHAPE.RECTANGLE, 0.15, 3.3, 3.2, 0.09, fill=GOLD)
    crown = c.shp(MSO_SHAPE.STAR_5_POINT, 1.5, 0.35, 0.55, 0.55, fill=GOLD, line=BROWN, lw=1.25)
    hw, hh = c.W(1.5), c.H(1.16)
    horse = horse_shape(s, c.X(1.0), c.Y(2.1), hw, hh, fill=BROWN)
    body = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 1.62, 1.62, 0.42, 0.62, fill=GREEN2)
    head = c.shp(MSO_SHAPE.OVAL, 1.68, 1.28, 0.32, 0.32, fill=TAN)
    flag = c.shp(MSO_SHAPE.RIGHT_TRIANGLE, 1.86, 0.95, 0.42, 0.4, fill=GREEN2, rot=0)
    lab = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 0.45, 3.65, 2.6, 0.55, fill=BROWN)
    _txt(lab, [("a Sultan, not a Sultana", 12.5, CREAM2, HEAD_FONT, False)])
    riders = [horse, body, head, flag]
    anim.add_group([{"spid": ground.shape_id, "kind": "wipe", "dur": 450},
                    {"spid": crown.shape_id, "kind": "zoom", "dur": 400, "delay": 300},
                    {"spid": lab.shape_id, "kind": "fade", "dur": 400, "delay": 600}],
                   click=False)
    anim.add_group([{"spid": sp.shape_id, "kind": "motion",
                     "path": (f"M -{c.fx(1.6):.3f} 0 L {c.fx(2.6):.3f} -0.008 "
                              f"L {c.fx(3.4):.3f} 0 E"),
                     "dur": 5200, "repeat": "indefinite"} for sp in riders],
                   click=False)
    anim.add_group([{"spid": crown.shape_id, "kind": "pulse", "dur": 900,
                     "repeat": "indefinite"}], click=False)
    for sp in [ground, crown, lab] + riders:
        anim.build_shape(sp.shape_id)


def scene_puppet(s, anim, L, T, W, H):
    """The Forty's puppet: strings and a swinging sultan."""
    c = Canvas(s, L, T, W, H)
    bar = c.shp(MSO_SHAPE.RECTANGLE, 0.55, 0.5, 2.4, 0.22, fill=DARK)
    forty = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 1.05, 0.05, 1.4, 0.42, fill=DARK, line=GOLD, lw=1.0)
    _txt(forty, [("THE FORTY", 11.5, CREAM2, HEAD_FONT, False)])
    s1 = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, c.X(1.1), c.Y(0.72),
                                c.X(1.5), c.Y(1.9))
    s2 = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, c.X(2.4), c.Y(0.72),
                                c.X(2.0), c.Y(1.9))
    for ln in (s1, s2):
        ln.line.color.rgb = GRAY; ln.line.width = Pt(1.25); ln.shadow.inherit = False
    pup = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 1.3, 1.9, 0.95, 1.25, fill=BROWN2, line=GOLD, lw=1.25)
    _txt(pup, [("PUPPET", 12, CREAM2, HEAD_FONT, False),
               ("SULTAN", 12, CREAM2, HEAD_FONT, False),
               ("Bahram \u2022 Masud", 9, WHITE, BODY_FONT, False)])
    lab = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 0.5, 3.6, 2.5, 0.5, fill=BROWN)
    _txt(lab, [("crowned, controlled, cut loose", 11, CREAM2, HEAD_FONT, False)])
    anim.add_group([{"spid": forty.shape_id, "kind": "zoom", "dur": 450},
                    {"spid": bar.shape_id, "kind": "wipe", "dur": 400, "delay": 350},
                    {"spid": s1.shape_id, "kind": "wipe", "dur": 350, "delay": 650},
                    {"spid": s2.shape_id, "kind": "wipe", "dur": 350, "delay": 800},
                    {"spid": pup.shape_id, "kind": "fly", "path": "M 0 -0.12 L 0 0 E",
                     "dur": 550, "delay": 1000},
                    {"spid": lab.shape_id, "kind": "fade", "dur": 400, "delay": 1500}],
                   click=False)
    anim.add_group([{"spid": pup.shape_id, "kind": "motion",
                     "path": (f"M 0 0 L {c.fx(0.22):.3f} {c.fy(0.1):.3f} L 0 {c.fy(0.16):.3f} "
                              f"L -{c.fx(0.22):.3f} {c.fy(0.1):.3f} L 0 0 E"),
                     "dur": 3000, "repeat": "indefinite"}], click=False)
    for sp in [bar, forty, s1, s2, pup, lab]:
        anim.build_shape(sp.shape_id)


def scene_balban(s, anim, L, T, W, H):
    """Two pillars under a golden beam — majesty and terror."""
    c = Canvas(s, L, T, W, H)
    p1 = c.shp(MSO_SHAPE.RECTANGLE, 0.75, 1.6, 0.85, 2.0, fill=BROWN, line=GOLD, lw=1.25)
    _txt(p1, [("THEATRE", 10.5, CREAM2, HEAD_FONT, False),
              ("OF MAJESTY", 10.5, CREAM2, HEAD_FONT, False),
              ("sijda \u2022 paibos", 8.5, WHITE, BODY_FONT, False)])
    p2 = c.shp(MSO_SHAPE.RECTANGLE, 1.9, 1.6, 0.85, 2.0, fill=BROWN2, line=GOLD, lw=1.25)
    _txt(p2, [("BLOOD", 10.5, CREAM2, HEAD_FONT, False),
              ("AND IRON", 10.5, CREAM2, HEAD_FONT, False),
              ("barid spies", 8.5, WHITE, BODY_FONT, False)])
    beam = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 0.5, 0.85, 2.5, 0.62, fill=GOLD, line=BROWN, lw=1.5)
    _txt(beam, [("ZIL-E-ILAHI \u2014 Shadow of God", 11.5, DARK, HEAD_FONT, True)])
    base = c.shp(MSO_SHAPE.RECTANGLE, 0.5, 3.6, 2.5, 0.16, fill=DARK)
    anim.add_group([{"spid": base.shape_id, "kind": "wipe", "dur": 400},
                    {"spid": p1.shape_id, "kind": "wipeup", "dur": 650, "delay": 300},
                    {"spid": p2.shape_id, "kind": "wipeup", "dur": 650, "delay": 800},
                    {"spid": beam.shape_id, "kind": "fly", "path": "M 0 -0.16 L 0 0 E",
                     "dur": 700, "delay": 1500}], click=False)
    anim.add_group([{"spid": beam.shape_id, "kind": "pulse", "dur": 900,
                     "repeat": "indefinite"}], click=False)
    for sp in [p1, p2, beam, base]:
        anim.build_shape(sp.shape_id)


def sword_clash_scene(s, anim, L, T, W, H, cap_text):
    """1290: two talwars fly in and cross; the clash star keeps ringing."""
    c = Canvas(s, L, T, W, H, cw=3.4, ch=2.9)
    sr = sword_shape(s, c.X(1.85), c.Y(0.55), c.W(0.34), c.H(2.0), mirror=False, rot=45)
    sl = sword_shape(s, c.X(1.2), c.Y(0.55), c.W(0.34), c.H(2.0), mirror=True, rot=-45)
    star = c.shp(MSO_SHAPE.STAR_8_POINT, 1.32, 1.0, 0.75, 0.75, fill=GOLD, line=BROWN, lw=1.25)
    cap = s.shapes.add_textbox(int(c.X(0.0)), int(c.Y(2.55)), int(c.W(3.4)), int(c.H(0.35)))
    p = cap.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = cap_text
    r.font.size = Pt(12); r.font.color.rgb = GRAY
    r.font.name = BODY_FONT; r.font.italic = True
    anim.add_group([
        {"spid": sr.shape_id, "kind": "fly", "path": "M 0.14 0.12 L 0 0 E",
         "dur": 600, "swing": 25},
        {"spid": sl.shape_id, "kind": "fly", "path": "M -0.14 0.12 L 0 0 E",
         "dur": 600, "swing": -25},
        {"spid": star.shape_id, "kind": "zoom", "dur": 300, "delay": 550},
        {"spid": cap.shape_id, "kind": "fade", "dur": 400, "delay": 800}], click=False)
    anim.add_group([{"spid": star.shape_id, "kind": "pulse", "dur": 700,
                     "repeat": "indefinite"}], click=False)
    for sp in (sr, sl, star, cap):
        anim.build_shape(sp.shape_id)


def scene_duel(s, anim, L, T, W, H):
    sword_clash_scene(s, anim, L, T, W, H, "1290 \u2014 the Khalji swords fall")


def scene_tanka(s, anim, L, T, W, H):
    """The silver tanka: a shine sweeps across, forever."""
    c = Canvas(s, L, T, W, H, cw=2.9, ch=2.9)
    coin = c.shp(MSO_SHAPE.OVAL, 0.45, 0.45, 2.0, 2.0, fill=SILVER, line=GOLD, lw=2.5)
    _txt(coin, [("tanka", 20, DARK, HEAD_FONT, True),
                ("silver \u2022 Iltutmish", 10.5, BROWN, BODY_FONT, False)])
    ring = c.shp(MSO_SHAPE.OVAL, 0.72, 0.72, 1.46, 1.46, fill=None, line=BROWN2, lw=1.25)
    shine = c.shp(MSO_SHAPE.RECTANGLE, 0.2, 0.3, 0.28, 2.3, fill=WHITE, rot=24)
    shine.fill.fore_color.rgb = WHITE
    lab = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 0.35, 2.6, 2.2, 0.42, fill=BROWN)
    _txt(lab, [("ancestor of the rupee", 11, CREAM2, HEAD_FONT, False)])
    anim.add_group([{"spid": coin.shape_id, "kind": "zoom", "dur": 550},
                    {"spid": ring.shape_id, "kind": "fade", "dur": 400, "delay": 450},
                    {"spid": lab.shape_id, "kind": "fade", "dur": 400, "delay": 750}],
                   click=False)
    anim.add_group([{"spid": shine.shape_id, "kind": "motion",
                     "path": f"M 0 0 L {c.fx(2.4):.3f} 0 E", "dur": 2600,
                     "repeat": "indefinite"}], click=False)
    for sp in (coin, ring, shine, lab):
        anim.build_shape(sp.shape_id)


def scene_minar(s, anim, L, T, W, H):
    """Qutb Minar builds itself storey by storey."""
    c = Canvas(s, L, T, W, H)
    ground = c.shp(MSO_SHAPE.RECTANGLE, 0.35, 3.85, 2.8, 0.1, fill=GOLD)
    storeys = [(1.25, 0.8, "1199"), (1.05, 0.72, ""), (0.86, 0.64, ""),
               (0.68, 0.54, "1220"), (0.52, 0.46, "")]
    blocks, y = [], 3.85
    for w, h, yr in storeys:
        y -= h
        b = c.shp(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, 1.75 - w / 2, y, w, h,
                  fill=BROWN if len(blocks) % 2 == 0 else BROWN2, line=GOLD, lw=1.0)
        if yr:
            _txt(b, [(yr, 9.5, CREAM2, BODY_FONT, True)])
        blocks.append(b)
    dome = c.shp(MSO_SHAPE.OVAL, 1.61, y - 0.26, 0.28, 0.26, fill=GOLD, line=BROWN, lw=1.0)
    lab = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 0.55, 0.2, 2.4, 0.5, fill=BROWN)
    _txt(lab, [("QUTB MINAR \u2014 72.5 m", 12.5, CREAM2, HEAD_FONT, False)])
    anim.add_group([{"spid": ground.shape_id, "kind": "wipe", "dur": 400}], click=False)
    effs = [{"spid": b.shape_id, "kind": "wipeup", "dur": 480, "delay": 300 + i * 480}
            for i, b in enumerate(blocks)]
    effs.append({"spid": dome.shape_id, "kind": "zoom", "dur": 350,
                 "delay": 300 + len(blocks) * 480})
    effs.append({"spid": lab.shape_id, "kind": "fade", "dur": 400,
                 "delay": 600 + len(blocks) * 480})
    anim.add_group(effs, click=False)
    for sp in blocks + [ground, dome, lab]:
        anim.build_shape(sp.shape_id)


def scene_legacy(s, anim, L, T, W, H):
    """What they left behind: chips radiate from the core."""
    c = Canvas(s, L, T, W, H)
    core = c.shp(MSO_SHAPE.OVAL, 0.95, 0.35, 1.6, 1.0, fill=DARK, line=GOLD, lw=1.5)
    _txt(core, [("1206\u20131290", 12, CREAM2, HEAD_FONT, False),
                ("WHAT THEY LEFT", 10.5, WHITE, BODY_FONT, True)])
    chips = [("DELHI as capital", BROWN), ("IQTA + ministries", BROWN2),
             ("TANKA currency", BROWN), ("Mongols kept out", BROWN2),
             ("Indo-Islamic culture", BROWN), ("a throne EARNED", GREEN2)]
    shapes, y = [], 1.62
    for txt, fill in chips:
        ch_ = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 0.55, y, 2.4, 0.42, fill=fill)
        _txt(ch_, [(txt, 10.5, WHITE, BODY_FONT, True)])
        shapes.append(ch_)
        y += 0.5
    anim.add_group([{"spid": core.shape_id, "kind": "zoom", "dur": 450}], click=False)
    anim.add_group([{"spid": ch_.shape_id, "kind": "fly",
                     "path": "M 0 -0.06 L 0 0 E", "dur": 350, "delay": 350 + i * 280}
                    for i, ch_ in enumerate(shapes)], click=False)
    anim.add_group([{"spid": core.shape_id, "kind": "pulse", "dur": 900,
                     "repeat": "indefinite"}], click=False)
    for sp in [core] + shapes:
        anim.build_shape(sp.shape_id)


# ================================================================= teddy
def teddy_scene(s, anim, L, T, W, H, caption):
    """The thank-you teddy, built from ovals — one arm waves forever."""
    c = Canvas(s, L, T, W, H, cw=4.2, ch=5.25)
    ear_l = c.shp(MSO_SHAPE.OVAL, 1.18, 0.55, 0.62, 0.62, fill=TEDDY)
    ear_r = c.shp(MSO_SHAPE.OVAL, 2.4, 0.55, 0.62, 0.62, fill=TEDDY)
    ein_l = c.shp(MSO_SHAPE.OVAL, 1.32, 0.69, 0.34, 0.34, fill=TEDDY2)
    ein_r = c.shp(MSO_SHAPE.OVAL, 2.54, 0.69, 0.34, 0.34, fill=TEDDY2)
    head = c.shp(MSO_SHAPE.OVAL, 1.28, 0.78, 1.64, 1.42, fill=TEDDY)
    muz = c.shp(MSO_SHAPE.OVAL, 1.72, 1.42, 0.76, 0.62, fill=TEDDY2)
    nose = c.shp(MSO_SHAPE.OVAL, 1.98, 1.5, 0.24, 0.18, fill=DARK)
    eye_l = c.shp(MSO_SHAPE.OVAL, 1.66, 1.16, 0.16, 0.18, fill=DARK)
    eye_r = c.shp(MSO_SHAPE.OVAL, 2.38, 1.16, 0.16, 0.18, fill=DARK)
    body = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 1.28, 2.18, 1.64, 1.72, fill=TEDDY)
    belly = c.shp(MSO_SHAPE.OVAL, 1.62, 2.5, 0.96, 1.05, fill=TEDDY2)
    arm_l = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 0.78, 2.35, 0.62, 1.15, fill=TEDDY, rot=18)
    arm_r = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 2.8, 2.15, 0.62, 1.15, fill=TEDDY, rot=-38)
    leg_l = c.shp(MSO_SHAPE.OVAL, 1.28, 3.78, 0.72, 0.6, fill=TEDDY, rot=12)
    leg_r = c.shp(MSO_SHAPE.OVAL, 2.2, 3.78, 0.72, 0.6, fill=TEDDY, rot=-12)
    heart = c.shp(MSO_SHAPE.HEART, 3.15, 1.05, 0.62, 0.58, fill=BROWN2)
    bub = c.shp(MSO_SHAPE.ROUNDED_RECTANGLE, 0.15, 0.1, 1.5, 0.62, fill=CREAM2, line=BROWN, lw=1.5)
    _txt(bub, [("thank you", 13, BROWN, HEAD_FONT, True),
               ("mates!", 13, BROWN, HEAD_FONT, True)])
    order = [leg_l, leg_r, body, belly, arm_l, arm_r, ear_l, ear_r, ein_l, ein_r,
             head, muz, nose, eye_l, eye_r, heart, bub]
    effs, d = [], 0
    for sp in order:
        effs.append({"spid": sp.shape_id, "kind": "fade", "dur": 160, "delay": d})
        d += 90
    anim.add_group(effs, click=False)
    # the wave: arm rocks back and forth, forever
    anim.add_group([{"spid": arm_r.shape_id, "kind": "spin", "deg": 46, "rev": True,
                     "dur": 800, "repeat": "indefinite", "delay": 1600}], click=False)
    # gentle happy bounce of the whole bear
    bounce = (f"M 0 0 L 0 -{c.fy(0.12):.3f} L 0 0 E")
    anim.add_group([{"spid": sp.shape_id, "kind": "motion", "path": bounce,
                     "dur": 1100, "delay": 1600, "repeat": "indefinite"}
                    for sp in (head, muz, nose, eye_l, eye_r, ear_l, ear_r,
                               ein_l, ein_r, body, belly)], click=False)
    anim.add_group([{"spid": heart.shape_id, "kind": "pulse", "dur": 800, "delay": 1600,
                     "repeat": "indefinite"}], click=False)
    cap = s.shapes.add_textbox(int(L - Inches(0.3)), int(T + H + Inches(0.06)),
                               int(W + Inches(0.6)), int(Inches(0.55)))
    p = cap.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = caption
    r.font.size = Pt(12); r.font.color.rgb = GRAY
    r.font.name = BODY_FONT; r.font.italic = True
    anim.add_group([{"spid": cap.shape_id, "kind": "fade", "dur": 500}], click=False)
    for sp in order + [cap]:
        anim.build_shape(sp.shape_id)
